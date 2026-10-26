"""Deliverable generator: Word (.docx) and PDF reports built from one audit-log row."""
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

TITLE = "Analysis & Recommendation Report"


def _fmt_time(iso) -> str:
    try:
        return datetime.fromisoformat(iso).strftime("%Y-%m-%d %H:%M UTC")
    except Exception:
        return iso or ""


def _sources(audit: dict) -> list[dict]:
    """Normalise sources: citation dicts, or plain strings from older/mock rows."""
    out = []
    for i, s in enumerate(audit.get("sources") or [], start=1):
        if isinstance(s, dict):
            out.append({"label": s.get("label", f"S{i}"), "source": s.get("source", ""),
                        "page": s.get("page", ""), "text": s.get("text", "")})
        else:
            out.append({"label": f"S{i}", "source": str(s), "page": "", "text": ""})
    return out


def _tools(audit: dict) -> list[dict]:
    out = []
    for t in audit.get("tools_used") or []:
        out.append(t if isinstance(t, dict) else {"tool": str(t), "input": "", "output": ""})
    return out


def _excerpt(text: str, n: int = 220) -> str:
    return text[:n] + "..." if len(text) > n else text


def build_docx(audit: dict, path: Path) -> Path:
    sources, tools = _sources(audit), _tools(audit)
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(TITLE, level=0)
    meta = doc.add_paragraph()
    meta.add_run(f"Report ID: {audit['id']}   |   Generated: {_fmt_time(audit.get('timestamp'))}   |   "
                 f"Requested by: {audit['user']}   |   Model: {audit.get('model') or ''}").italic = True

    doc.add_heading("Question", level=1)
    doc.add_paragraph(audit["question"])
    doc.add_heading("Findings", level=1)
    doc.add_paragraph(audit.get("answer") or "")

    calcs = [t for t in tools if t.get("tool") == "calculator"]
    if calcs:
        doc.add_heading("Calculations", level=1)
        for c in calcs:
            doc.add_paragraph(f"{c.get('input', '')}  ->  {c.get('output', '')}", style="List Bullet")

    doc.add_heading("Sources", level=1)
    if sources:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        for cell, head in zip(table.rows[0].cells, ["Ref", "Document", "Page", "Excerpt"]):
            cell.text = head
        for s in sources:
            row = table.add_row().cells
            row[0].text, row[1].text, row[2].text = s["label"], s["source"], str(s["page"])
            row[3].text = _excerpt(s["text"])
    else:
        doc.add_paragraph("No supporting sources: the answer was not found in the provided documents.")

    doc.add_heading("Tools used", level=1)
    for t in tools:
        doc.add_paragraph(f"{t.get('tool')}: {t.get('output', '')}", style="List Bullet")
    foot = doc.add_paragraph()
    foot.add_run("Generated locally by ENCLAVE AI. No data left this machine.").italic = True
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def build_pdf(audit: dict, path: Path) -> Path:
    sources, tools = _sources(audit), _tools(audit)
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=10.5, leading=15)
    small = ParagraphStyle("small", parent=body, fontSize=8.5, leading=11, textColor=colors.HexColor("#444444"))
    p = lambda text, style=body: Paragraph(escape(str(text)).replace("\n", "<br/>"), style)

    story = [Paragraph(TITLE, styles["Title"]),
             p(f"Report ID {audit['id']}  |  {_fmt_time(audit.get('timestamp'))}  |  "
               f"Requested by {audit['user']}  |  Model {audit.get('model') or ''}", small),
             Spacer(1, 12), Paragraph("Question", styles["Heading2"]), p(audit["question"]),
             Paragraph("Findings", styles["Heading2"]), p(audit.get("answer") or "")]

    calcs = [t for t in tools if t.get("tool") == "calculator"]
    if calcs:
        story.append(Paragraph("Calculations", styles["Heading2"]))
        story += [p(f"- {c.get('input', '')}  ->  {c.get('output', '')}") for c in calcs]

    story.append(Paragraph("Sources", styles["Heading2"]))
    if sources:
        rows = [[p("Ref", small), p("Document", small), p("Page", small), p("Excerpt", small)]]
        for s in sources:
            rows.append([p(s["label"], small), p(s["source"], small), p(s["page"], small),
                         p(_excerpt(s["text"]), small)])
        t = Table(rows, colWidths=[1.2 * cm, 6.3 * cm, 1.3 * cm, 8.2 * cm], repeatRows=1)
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                               ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef7")),
                               ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(t)
    else:
        story.append(p("No supporting sources: the answer was not found in the provided documents."))

    story += [Spacer(1, 14), p("Generated locally by ENCLAVE AI. No data left this machine.", small)]
    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                      topMargin=2 * cm, bottomMargin=2 * cm, title=TITLE).build(story)
    return path
